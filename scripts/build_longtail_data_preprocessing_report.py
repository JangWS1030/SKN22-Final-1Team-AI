from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Dict, Iterable, List

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


COMMUNITY_PRICE = 0.34
SECURE_PRICE = 0.59
NETWORK_VOLUME_PER_GB_MONTH = 0.07


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a DOCX preprocessing report for long-tail hairstyle training data."
    )
    parser.add_argument(
        "--catalog",
        type=Path,
        default=Path("dataset_build/configs/longtail_hair_sources.json"),
    )
    parser.add_argument(
        "--download-report",
        type=Path,
        default=Path("dataset_build/raw/public_hair_sources/download_report.json"),
    )
    parser.add_argument(
        "--face-sketches-summary",
        type=Path,
        default=Path("dataset_build/processed/face_sketches_refined_generation/reports/summary.json"),
    )
    parser.add_argument(
        "--male-asian-summary",
        type=Path,
        default=Path("dataset_build/processed/male_asian_hairstyles_generation/reports/summary.json"),
    )
    parser.add_argument(
        "--subset-summary",
        type=Path,
        default=Path("dataset_build/processed/longtail_training/reports/longtail_training_summary.json"),
    )
    parser.add_argument(
        "--base-summary",
        type=Path,
        default=Path("dataset_build/processed/celeba_dialog_hq_generation/reports/style_training_subsets_summary.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/MirrAI_데이터 전처리_장기학습용 인공지능 데이터 전처리 결과서.docx"),
    )
    return parser.parse_args()


def load_json(path: Path) -> Dict:
    return json.loads(path.read_text(encoding="utf-8"))


def dir_size_gb(path: Path) -> float:
    if not path.exists():
        return 0.0
    size = sum(p.stat().st_size for p in path.rglob("*") if p.is_file())
    return round(size / (1024 ** 3), 3)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:color"), "999999")


def set_run_font(run, size_pt: float, *, bold: bool = False, color: str | None = None) -> None:
    run.bold = bold
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, 10.5)


def add_heading(doc: Document, text: str, *, major: bool = True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if major else 8)
    p.paragraph_format.space_after = Pt(6 if major else 4)
    run = p.add_run(text)
    set_run_font(run, 13 if major else 11, bold=True)


def add_table(doc: Document, headers: List[str], rows: Iterable[Iterable[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_borders(table)
    hdr = table.rows[0]
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.text = value
        set_cell_shading(cell, "666666")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, 10.5, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, 10.5, bold=(idx == 0))
    return table


def add_code_block_table(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders(table)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F7F7F7")
    cell.text = text
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            set_run_font(run, 10.0)
    return table


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.18)
    section.bottom_margin = Inches(1.00)
    section.left_margin = Inches(1.00)
    section.right_margin = Inches(1.00)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10.5)


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("SK네트웍스 Family AI과정 22기\n")
    set_run_font(run, 13)
    gray = p.add_run("데이터 전처리 ")
    set_run_font(gray, 18, bold=True, color="7F7F7F")
    title = p.add_run("인공지능 데이터 전처리 결과서")
    set_run_font(title, 18, bold=True)


def build_budget_rows(total_storage_gb: float) -> List[List[str]]:
    recommended_hours = 22.0
    extended_hours = 30.0
    volume_gb = max(80.0, round(total_storage_gb + 20.0, 0))
    volume_cost = round(volume_gb * NETWORK_VOLUME_PER_GB_MONTH, 2)
    recommended_community = round(recommended_hours * COMMUNITY_PRICE, 2)
    recommended_secure = round(recommended_hours * SECURE_PRICE + volume_cost, 2)
    extended_community = round(extended_hours * COMMUNITY_PRICE, 2)
    extended_secure = round(extended_hours * SECURE_PRICE + volume_cost, 2)
    return [
        [
            "권장 예산",
            f"4090 {recommended_hours:.0f}시간 + {volume_gb:.0f}GB 볼륨",
            f"${recommended_community:.2f} (GPU만)",
            f"${recommended_secure:.2f} (GPU+볼륨)",
        ],
        [
            "여유 예산",
            f"4090 {extended_hours:.0f}시간 + {volume_gb:.0f}GB 볼륨",
            f"${extended_community:.2f} (GPU만)",
            f"${extended_secure:.2f} (GPU+볼륨)",
        ],
        [
            "비고",
            "stage3 장기 스타일 + stage4 garment reveal + 1~2회 재시도 기준",
            "$0.34/hr",
            f"$0.59/hr + ${volume_cost:.2f}/mo",
        ],
    ]


def main() -> None:
    args = parse_args()
    catalog = load_json(args.catalog)
    download_report = load_json(args.download_report)
    face_sketches = load_json(args.face_sketches_summary)
    male_asian = load_json(args.male_asian_summary)
    subset_summary = load_json(args.subset_summary)
    base_summary = load_json(args.base_summary)

    current_raw_gb = dir_size_gb(Path("dataset_build/raw/celeba_dialog_hq"))
    current_processed_gb = dir_size_gb(Path("dataset_build/processed/celeba_dialog_hq_generation"))
    public_raw_gb = dir_size_gb(Path("dataset_build/raw/public_hair_sources"))
    public_processed_gb = (
        dir_size_gb(Path("dataset_build/processed/face_sketches_refined_generation"))
        + dir_size_gb(Path("dataset_build/processed/male_asian_hairstyles_generation"))
        + dir_size_gb(Path("dataset_build/processed/longtail_training"))
    )
    total_storage_gb = round(current_raw_gb + current_processed_gb + public_raw_gb + public_processed_gb, 3)

    doc = Document()
    configure_document(doc)
    add_title_block(doc)

    add_heading(doc, "1. 문서 개요")
    add_body_paragraph(doc, "프로젝트명: MirrAI - 장기학습용 long-tail 헤어스타일 및 garment reveal 학습 데이터셋 구축")
    add_body_paragraph(doc, "전처리 목적: 기존 CelebA-Dialog 기반 생성 학습셋에 비주류 헤어스타일과 긴머리 제거 후 의복/목선 복원용 데이터를 추가해 RunPod 4090 이어학습 품질을 높인다.")
    add_body_paragraph(doc, "문제 정의: 기존 파이프라인은 일반 short/medium 스타일에는 강하지만 afro, cornrows, dreadlocks, liberty spikes, hime cut 같은 long-tail 스타일과 long-to-bob 전환 시 가려진 옷 복원 능력이 약하다.")
    add_body_paragraph(doc, "추가 고려 사항: 공개 소스의 라이선스, 스타일 라벨 신뢰도, 실제 학습 manifest 호환성, RunPod Pod 저장공간 비용을 함께 고려해 자동 다운로드 가능 소스와 수동 반입 소스를 분리한다.")

    add_heading(doc, "2. 데이터셋 개요")
    dataset_rows = [
        ["기존 베이스", "CelebA-Dialog HQ 전처리셋", f"학습용 {base_summary['train_highconf']}건", "현재 stage2 LoRA 기반 안정화용"],
        ["추가 long-tail", "FaceSketches-HairStyle40 Refined", f"원본 204장 / 수용 {face_sketches['accepted_recon_samples']}건", "클래스 라벨이 뚜렷한 rare style 보강"],
        ["옵션 support", "Male Asian Hairstyles", f"원본 65장 / 수용 {male_asian['accepted_recon_samples']}건", "기본 학습 제외, 참고용/OOD 확인용"],
        ["메타데이터", "Google Open Images Hair Style", "메타데이터만 반입", "후속 대규모 크롤링 후보 탐색용"],
        ["수동 권장", "FairFace / Figaro1k / LaPa / Openverse", "별도 반입", "균형 regularizer, rare texture, mask 품질, 실사용 long-tail 수집"],
    ]
    add_body_paragraph(doc, "데이터 출처 및 수집 방법: 허깅페이스 공개 소스는 자동 다운로드 스크립트로 반입하고, 라이선스/규모 이슈가 있는 소스는 별도 수동 반입 대상으로 catalog에 등록했다.")
    add_body_paragraph(doc, "최종 학습 직접 사용 범위: stage3는 기존 CelebA 고신뢰 셋 + FaceSketches refined external self-reconstruction 셋을 혼합하고, stage4는 garment reveal 후보 서브셋을 별도로 사용한다. male_asian_hairstyles는 기본 train mix에서 제외한다.")
    add_table(doc, ["구분", "데이터셋", "사용 범위", "비고"], dataset_rows)

    add_heading(doc, "3. 전처리 프로세스 개요")
    add_body_paragraph(doc, "전체 흐름도: 공개 소스 다운로드 -> 이미지 폴더 스캔 -> BiSeNet face parsing으로 hair/face/cloth mask 생성 -> 512 정규화 -> canny/control/face crop 생성 -> self-reconstruction manifest 저장 -> stage3/stage4 서브셋 분리")
    process_rows = [
        ["소스 수집", "공개 long-tail 확보", "HF snapshot 다운로드, manual source catalog 작성", "huggingface_hub, JSON catalog"],
        ["라벨 정규화", "스타일명을 prompt에 연결", "folder label humanize, family mapping, rarity weight 부여", "regex, alias map"],
        ["마스크 생성", "학습 호환 asset 생성", "BiSeNet으로 hair/face/cloth parsing 후 PNG 저장", "torch, face parsing checkpoint"],
        ["제어 입력 생성", "현재 학습기 입력 맞춤", "canny control image, face crop, 512 resize", "opencv, PIL"],
        ["subset 분리", "학습 단계별 manifest 생성", "stage3 rare style mix / stage4 garment reveal 분리", "JSONL builder"],
    ]
    add_table(doc, ["단계", "목적", "수행 작업", "사용 도구/라이브러리"], process_rows)

    add_heading(doc, "4. 세부 전처리 단계")
    add_heading(doc, "4.1 결측치 및 다운로드 정책", major=False)
    add_body_paragraph(doc, "자동 다운로드 대상은 HF 공개 소스 3종으로 제한하고, FairFace·Figaro1k·LaPa·FFHQ·Openverse는 라이선스/규모/품질 검수 이유로 catalog에만 등록했다.")
    add_heading(doc, "4.2 이상치 및 노이즈 처리", major=False)
    filter_rows = [
        ["얼굴 비율", "0.04 이상 0.68 이하", "얼굴 검출 실패 또는 너무 작은 샘플 제외"],
        ["머리 비율", "0.012 이상 0.70 이하", "머리 없음 또는 과도한 오버세그멘테이션 제외"],
        ["모자 비율", "0.08 초과 제외", "hat 가림이 큰 샘플 제거"],
        ["garment 후보", "긴 스타일 또는 cloth_ratio 0.06 이상", "long-to-short 의복 복원 stage4 후보로 분리"],
    ]
    add_table(doc, ["구분", "판단 기준", "처리 방법"], filter_rows)
    add_heading(doc, "4.3 데이터 변환 및 생성", major=False)
    add_body_paragraph(doc, "입력 형식: 공개 이미지 폴더 구조(core/reference_only/holdout 또는 long/medium/short 폴더)")
    add_body_paragraph(doc, "출력 형식: images_512/*.png, masks/hair/*.png, masks/face_protect/*.png, masks/cloth_protect/*.png, controls/canny/*.png, face_crops/*.png, manifests/*.jsonl")
    add_body_paragraph(doc, "주요 스크립트: download_public_hair_datasets.py, preprocess_external_hair_generation.py, build_longtail_training_subsets.py")
    command_text = "\n".join(
        [
            "주요 실행 명령",
            "python scripts/download_public_hair_datasets.py --include-optional",
            "docker run --rm -e PYTHONPATH=/workspace/repo -v <repo>:/workspace/repo --workdir /workspace/repo --entrypoint python byoungj/sd:llm-refined-trends-20260327-094231 scripts/preprocess_external_hair_generation.py --source-root dataset_build/raw/public_hair_sources/hf/face_sketches_refined --out-root dataset_build/processed/face_sketches_refined_generation --source-name face_sketches_refined --device cpu",
            "docker run --rm -e PYTHONPATH=/workspace/repo -v <repo>:/workspace/repo --workdir /workspace/repo --entrypoint python byoungj/sd:llm-refined-trends-20260327-094231 scripts/preprocess_external_hair_generation.py --source-root dataset_build/raw/public_hair_sources/hf/male_asian_hairstyles --out-root dataset_build/processed/male_asian_hairstyles_generation --source-name male_asian_hairstyles --device cpu",
            "python scripts/build_longtail_training_subsets.py",
        ]
    )
    add_code_block_table(doc, command_text)
    add_heading(doc, "4.4 학습/검증 분리", major=False)
    add_body_paragraph(doc, "FaceSketches refined는 core=train, reference_only=val, holdout=test 성격을 우선 존중했고, 기타 폴더 구조가 없는 소스는 hash split으로 train/val/test를 나눴다.")
    add_body_paragraph(doc, "stage3는 rare-style 유지가 목적이므로 base stability mix를 유지하되 external rows의 sample_weight를 더 높였고, stage4는 cloth/neckline이 드러나는 long-hair 후보만 따로 추렸다.")

    add_heading(doc, "5. 전처리 결과 요약 및 학습 예산")
    result_rows = [
        ["기존 베이스", f"raw {current_raw_gb:.3f}GB / processed {current_processed_gb:.3f}GB", f"train_highconf {base_summary['train_highconf']}건", "stage2 안정화 베이스"],
        ["FaceSketches", f"raw 포함 {public_raw_gb:.3f}GB", f"수용 {face_sketches['accepted_recon_samples']}건", "rare style class 보강"],
        ["Male Asian", f"processed {dir_size_gb(Path('dataset_build/processed/male_asian_hairstyles_generation')):.3f}GB", f"수용 {male_asian['accepted_recon_samples']}건", "기본 train 제외, 참고용 유지"],
        [
            "Stage3 mix",
            f"총 {subset_summary['stage3_longtail_mix_rows']}건",
            f"external {subset_summary['stage3_longtail_source_counts'].get('face_sketches_refined', 0) + subset_summary['stage3_longtail_source_counts'].get('male_asian_hairstyles', 0)}건",
            "rare style continuation",
        ],
        [
            "Stage4 garment",
            f"총 {subset_summary['stage4_garment_reveal_rows']}건",
            f"external {subset_summary['stage4_garment_source_counts'].get('face_sketches_refined', 0) + subset_summary['stage4_garment_source_counts'].get('male_asian_hairstyles', 0)}건",
            "long-to-short garment reveal",
        ],
    ]
    add_table(doc, ["구분", "전처리 결과", "직접 학습 사용 범위", "효과"], result_rows)
    add_body_paragraph(doc, f"권장 저장공간: 현재 확인된 raw+processed footprint는 약 {total_storage_gb:.3f}GB이며, 체크포인트/로그/여유 공간을 포함하면 80GB Secure volume 또는 64GB 이상 로컬 디스크를 권장한다.")
    add_body_paragraph(doc, "RunPod 4090 예산 산정은 2026-03-27 기준 공식 단가 Community $0.34/hr, Secure $0.59/hr를 사용했고, Secure Pod persistent volume은 $0.07/GB/month 기준으로 계산했다.")
    add_table(doc, ["구분", "예상 소요 시간", "Community 4090", "Secure 4090"], build_budget_rows(total_storage_gb))

    add_heading(doc, "6. 향후 사용 방안 및 참고 경로")
    add_body_paragraph(doc, "RunPod 4090 학습 순서: stage3_longtail_mix로 rare style continuation -> qualitative rare-style 평가 -> stage4_garment_reveal로 long-to-bob 의복/목선 복원 보강 -> endpoint 재검증")
    add_body_paragraph(doc, "추천 운영 방안: FairFace는 균형 regularizer, Figaro1k는 rare texture oversampling, Openverse/Wikimedia Commons는 truly long-tail named style 수집 경로로 별도 유지한다. male_asian_hairstyles는 품질/성별 편향 점검용 참고셋으로만 사용한다.")
    add_body_paragraph(doc, "참고 경로: dataset_build/configs/longtail_hair_sources.json, dataset_build/raw/public_hair_sources/download_report.json, dataset_build/processed/face_sketches_refined_generation/reports/summary.json, dataset_build/processed/male_asian_hairstyles_generation/reports/summary.json, dataset_build/processed/longtail_training/reports/longtail_training_summary.json")
    add_body_paragraph(doc, "외부 참고: FairFace(https://github.com/dchen236/FairFace), Figaro1k(https://www.michelesvanera.org/figaro-1k/), LaPa(https://github.com/jd-opensource/lapa-dataset), FFHQ(https://github.com/NVlabs/ffhq-dataset), Openverse docs(https://docs.openverse.org/index.html), RunPod pricing(https://www.runpod.io/pricing), RunPod network volumes(https://docs.runpod.io/storage/network-volumes)")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(args.output))
    print(json.dumps({"output": str(args.output.as_posix())}, ensure_ascii=False))


if __name__ == "__main__":
    main()
