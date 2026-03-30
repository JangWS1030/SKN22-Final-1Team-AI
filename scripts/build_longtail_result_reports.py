from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build DOCX reports for long-tail training and benchmark results.")
    parser.add_argument("--stage3-summary", type=Path, required=True)
    parser.add_argument("--stage4-summary", type=Path, default=None)
    parser.add_argument("--smoke-summary", type=Path, required=True)
    parser.add_argument("--rare-summary", type=Path, required=True)
    parser.add_argument("--garment-summary", type=Path, required=True)
    parser.add_argument("--smoke-predictions", type=Path, default=None)
    parser.add_argument("--rare-predictions", type=Path, default=None)
    parser.add_argument("--garment-predictions", type=Path, default=None)
    parser.add_argument("--work-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def load_json(path: Optional[Path]) -> Dict:
    if path is None or not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8-sig"))


def load_jsonl(path: Optional[Path]) -> List[Dict]:
    if path is None or not path.exists():
        return []
    rows: List[Dict] = []
    with path.open("r", encoding="utf-8-sig") as handle:
        for line in handle:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


def set_run_font(run, size_pt: float, *, bold: bool = False, color: str | None = None) -> None:
    run.bold = bold
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:color"), "999999")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.18)
    section.bottom_margin = Inches(1.00)
    section.left_margin = Inches(1.00)
    section.right_margin = Inches(1.00)
    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10.5)


def add_title_block(doc: Document, subtitle: str, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("SKN22 Final Project\n")
    set_run_font(run, 13)
    gray = p.add_run(subtitle)
    set_run_font(gray, 18, bold=True, color="7F7F7F")
    black = p.add_run(title)
    set_run_font(black, 18, bold=True)


def add_heading(doc: Document, text: str, *, major: bool = True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if major else 8)
    p.paragraph_format.space_after = Pt(6 if major else 4)
    run = p.add_run(text)
    set_run_font(run, 13 if major else 11, bold=True)


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, 10.5)


def add_table(doc: Document, headers: List[str], rows: Iterable[Iterable[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_borders(table)
    header_row = table.rows[0]
    for idx, value in enumerate(headers):
        cell = header_row.cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.text = value
        set_cell_shading(cell, "666666")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, 10.5, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cells[idx].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(run, 10.5, bold=(idx == 0))
    return table


def add_image_to_cell(cell, image_path: Optional[Path], width_inch: float = 1.75) -> None:
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if image_path is None or not image_path.exists():
        cell.text = "이미지 없음"
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, 9.5)
        return
    run = cell.paragraphs[0].add_run()
    run.add_picture(str(image_path), width=Inches(width_inch))


def fmt_float(value: object, digits: int = 4) -> str:
    if value is None:
        return "-"
    try:
        return f"{float(value):.{digits}f}"
    except (TypeError, ValueError):
        return str(value)


def fmt_int(value: object) -> str:
    if value is None:
        return "-"
    try:
        return f"{int(value)}"
    except (TypeError, ValueError):
        return str(value)


def fmt_bool(value: object) -> str:
    if value is None:
        return "-"
    return "예" if bool(value) else "아니오"


def metric_mean(summary: Dict, name: str, digits: int = 4) -> str:
    metrics = summary.get("metrics") or {}
    metric = metrics.get(name) or {}
    return fmt_float(metric.get("mean"), digits=digits)


def output_row(stage_name: str, summary: Dict) -> List[str]:
    config = summary.get("config") or {}
    return [
        stage_name,
        fmt_int(summary.get("global_step")),
        fmt_int(summary.get("best_step")),
        fmt_float(summary.get("best_val_loss")),
        fmt_bool(summary.get("stopped_early")),
        fmt_int(config.get("max_train_steps")),
        fmt_int(config.get("max_train_samples")),
    ]


def select_examples(rows: List[Dict], limit: int) -> List[Dict]:
    examples: List[Dict] = []
    for row in rows:
        prediction_path = Path(str(row.get("prediction_image_path") or ""))
        if not prediction_path.exists():
            continue
        examples.append(row)
        if len(examples) >= limit:
            break
    return examples


def benchmark_label(row: Dict) -> str:
    parts = [
        str(row.get("sample_id") or "").strip(),
        str(row.get("hairstyle_text") or row.get("hairstyle_id") or "").strip(),
    ]
    return " / ".join([part for part in parts if part])


def add_example_gallery(doc: Document, title: str, rows: List[Dict], limit: int) -> None:
    add_heading(doc, title, major=False)
    examples = select_examples(rows, limit)
    if not examples:
        add_body_paragraph(doc, "표시할 예시 이미지가 아직 없습니다.")
        return

    for row in examples:
        add_body_paragraph(doc, benchmark_label(row))
        table = doc.add_table(rows=2, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.style = "Table Grid"
        set_table_borders(table)

        headers = ["입력 이미지", "생성 결과", "정답 이미지"]
        for idx, header in enumerate(headers):
            cell = table.rows[0].cells[idx]
            cell.text = header
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, "666666")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, 10.0, bold=True, color="FFFFFF")

        source_path = Path(str(row.get("source_image_path") or ""))
        prediction_path = Path(str(row.get("prediction_image_path") or ""))
        target_path = Path(str(row.get("target_image_path") or ""))
        add_image_to_cell(table.rows[1].cells[0], source_path)
        add_image_to_cell(table.rows[1].cells[1], prediction_path)
        add_image_to_cell(table.rows[1].cells[2], target_path)


def add_training_report(
    path: Path,
    work_root: Path,
    stage3_summary: Dict,
    stage4_summary: Dict,
    promoted_lora: str,
    smoke_predictions: List[Dict],
) -> None:
    doc = Document()
    configure_document(doc)
    add_title_block(doc, "비주류 헤어 학습 ", "결과서")

    add_heading(doc, "1. 개요")
    add_body_paragraph(
        doc,
        "본 문서는 long-tail 헤어스타일 보강 학습(stage3)과 긴머리 제거 후 의상 및 목선 복원 학습(stage4)의 실행 결과를 정리한 학습 결과서이다.",
    )
    add_table(
        doc,
        ["항목", "내용"],
        [
            ["작업 루트", str(work_root)],
            ["stage3 요약", str(work_root / "output" / "training" / "generation_lora_stage3_longtail_4090" / "training_summary.json")],
            ["stage4 요약", str(work_root / "output" / "training" / "generation_lora_stage4_garment_reveal_4090" / "training_summary.json") if stage4_summary else "미실행"],
            ["최종 적용 LoRA", promoted_lora],
            ["문서 생성 시각", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ],
    )

    add_heading(doc, "2. 단계별 학습 결과")
    stage_rows = [output_row("stage3 longtail", stage3_summary)]
    if stage4_summary:
        stage_rows.append(output_row("stage4 garment reveal", stage4_summary))
    add_table(
        doc,
        ["단계", "종료 step", "best step", "best val loss", "조기 종료", "max step", "학습 샘플 수"],
        stage_rows,
    )

    add_heading(doc, "3. 주요 학습 설정", major=False)
    config = stage4_summary.get("config") or stage3_summary.get("config") or {}
    add_table(
        doc,
        ["항목", "값"],
        [
            ["해상도", fmt_int(config.get("resolution"))],
            ["batch size", fmt_int(config.get("train_batch_size"))],
            ["gradient accumulation", fmt_int(config.get("gradient_accumulation_steps"))],
            ["LoRA rank", fmt_int(config.get("rank"))],
            ["checkpoint 간격", fmt_int(config.get("checkpointing_steps"))],
            ["validation 간격", fmt_int(config.get("validation_steps"))],
            ["초기 LoRA", str(config.get("initial_lora_path") or "-")],
        ],
    )

    add_heading(doc, "4. 대표 생성 예시", major=False)
    add_body_paragraph(doc, "학습 완료 후 smoke benchmark를 기준으로 대표 생성 결과를 정리하였다.")
    add_example_gallery(doc, "Smoke benchmark 예시", smoke_predictions, limit=2)

    add_heading(doc, "5. 경로 및 산출물", major=False)
    add_table(
        doc,
        ["항목", "경로"],
        [
            ["stage3 output", str(work_root / "output" / "training" / "generation_lora_stage3_longtail_4090")],
            ["stage4 output", str(work_root / "output" / "training" / "generation_lora_stage4_garment_reveal_4090")],
            ["로그", str(work_root / "logs_longtail")],
            ["상태", str(work_root / "state_longtail")],
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def add_test_report(
    path: Path,
    work_root: Path,
    promoted_lora: str,
    smoke_summary: Dict,
    rare_summary: Dict,
    garment_summary: Dict,
    smoke_predictions: List[Dict],
    rare_predictions: List[Dict],
    garment_predictions: List[Dict],
) -> None:
    doc = Document()
    configure_document(doc)
    add_title_block(doc, "비주류 헤어 테스트 ", "결과서")

    add_heading(doc, "1. 개요")
    add_body_paragraph(
        doc,
        "본 문서는 학습 완료 후 smoke, rare-style showcase, garment reveal showcase 벤치마크를 실행한 결과를 정리한 테스트 결과서이다.",
    )
    add_table(
        doc,
        ["항목", "내용"],
        [
            ["적용 LoRA", promoted_lora],
            ["벤치마크 루트", str(work_root / "output" / "benchmarks")],
            ["문서 생성 시각", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ],
    )

    add_heading(doc, "2. 벤치마크 요약")
    add_table(
        doc,
        ["벤치마크", "샘플 수", "hair MAE", "hair PSNR", "hair SSIM", "face preserve MAE", "cloth preserve MAE"],
        [
            [
                "smoke",
                fmt_int(smoke_summary.get("rows")),
                metric_mean(smoke_summary, "hair_mae"),
                metric_mean(smoke_summary, "hair_psnr", digits=2),
                metric_mean(smoke_summary, "hair_ssim"),
                metric_mean(smoke_summary, "face_preserve_mae"),
                metric_mean(smoke_summary, "cloth_preserve_mae"),
            ],
            [
                "rare_style_showcase",
                fmt_int(rare_summary.get("rows")),
                metric_mean(rare_summary, "hair_mae"),
                metric_mean(rare_summary, "hair_psnr", digits=2),
                metric_mean(rare_summary, "hair_ssim"),
                metric_mean(rare_summary, "face_preserve_mae"),
                metric_mean(rare_summary, "cloth_preserve_mae"),
            ],
            [
                "garment_reveal_showcase",
                fmt_int(garment_summary.get("rows")),
                metric_mean(garment_summary, "hair_mae"),
                metric_mean(garment_summary, "hair_psnr", digits=2),
                metric_mean(garment_summary, "hair_ssim"),
                metric_mean(garment_summary, "face_preserve_mae"),
                metric_mean(garment_summary, "cloth_preserve_mae"),
            ],
        ],
    )

    add_heading(doc, "3. 예시 이미지", major=False)
    add_example_gallery(doc, "Smoke benchmark 예시", smoke_predictions, limit=1)
    add_example_gallery(doc, "Rare-style showcase 예시", rare_predictions, limit=2)
    add_example_gallery(doc, "Garment reveal showcase 예시", garment_predictions, limit=2)

    add_heading(doc, "4. 결과 경로", major=False)
    add_table(
        doc,
        ["항목", "경로"],
        [
            ["smoke report", str(work_root / "output" / "benchmarks" / "longtail_smoke" / "report" / "README.md")],
            ["rare-style report", str(work_root / "output" / "benchmarks" / "rare_style_showcase" / "report" / "README.md")],
            ["garment reveal report", str(work_root / "output" / "benchmarks" / "garment_reveal_showcase" / "report" / "README.md")],
            ["예측 이미지 루트", str(work_root / "output" / "benchmarks")],
        ],
    )

    add_heading(doc, "5. 해석 메모", major=False)
    add_body_paragraph(doc, "hair MAE는 낮을수록 좋고, hair PSNR 및 hair SSIM은 높을수록 좋다.")
    add_body_paragraph(doc, "face preserve MAE와 cloth preserve MAE는 값이 낮을수록 원본 얼굴 및 의상 보존이 양호하다는 의미이다.")
    add_body_paragraph(doc, "rare_style_showcase는 비주류 스타일 재현성, garment_reveal_showcase는 긴머리 제거 후 목선 및 옷 복원 품질을 중점적으로 확인한다.")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def resolve_promoted_lora(work_root: Path) -> str:
    candidates = [
        work_root / "output" / "training" / "generation_lora_stage4_garment_reveal_4090" / "best",
        work_root / "output" / "training" / "generation_lora_stage4_garment_reveal_4090" / "final",
        work_root / "output" / "training" / "generation_lora_stage3_longtail_4090" / "best",
        work_root / "output" / "training" / "generation_lora_stage3_longtail_4090" / "final",
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    return str(candidates[-1])


def main() -> None:
    args = parse_args()
    output_dir = args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)

    stage3_summary = load_json(args.stage3_summary)
    stage4_summary = load_json(args.stage4_summary)
    smoke_summary = load_json(args.smoke_summary)
    rare_summary = load_json(args.rare_summary)
    garment_summary = load_json(args.garment_summary)
    smoke_predictions = load_jsonl(args.smoke_predictions)
    rare_predictions = load_jsonl(args.rare_predictions)
    garment_predictions = load_jsonl(args.garment_predictions)

    promoted_lora = resolve_promoted_lora(args.work_root)

    training_report_path = output_dir / "MirrAI_비주류_헤어_학습_결과서.docx"
    test_report_path = output_dir / "MirrAI_비주류_헤어_테스트_결과서.docx"

    add_training_report(
        training_report_path,
        args.work_root,
        stage3_summary,
        stage4_summary,
        promoted_lora,
        smoke_predictions,
    )
    add_test_report(
        test_report_path,
        args.work_root,
        promoted_lora,
        smoke_summary,
        rare_summary,
        garment_summary,
        smoke_predictions,
        rare_predictions,
        garment_predictions,
    )

    summary = {
        "training_report_path": str(training_report_path),
        "test_report_path": str(test_report_path),
        "promoted_lora": promoted_lora,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }
    (output_dir / "report_bundle_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
