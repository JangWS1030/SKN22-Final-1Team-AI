from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

from docx import Document
from docx.enum.text import WD_BREAK


REPO_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = REPO_ROOT / "docs"
OUTPUT_DIR = REPO_ROOT / "output" / "doc"
REPORTS_DIR = REPO_ROOT / "dataset_build" / "processed" / "celeba_dialog_hq_generation" / "reports"
CONFIG_DIR = REPO_ROOT / "configs" / "generation_train"


def load_json(path: Path) -> Dict:
    return json.loads(path.read_text(encoding="utf-8"))


def format_bytes(num_bytes: int) -> str:
    units = ["B", "KB", "MB", "GB", "TB"]
    value = float(num_bytes)
    unit = units[0]
    for unit in units:
        if value < 1024.0 or unit == units[-1]:
            break
        value /= 1024.0
    return f"{value:.2f} {unit}"


def compute_dir_size(path: Path) -> int:
    return sum(file.stat().st_size for file in path.rglob("*") if file.is_file())


def replace_paragraph_prefixes(doc: Document, replacements: Dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        for prefix, value in replacements.items():
            if text.startswith(prefix):
                paragraph.text = value
                break


def add_table(document: Document, headers: Iterable[str], rows: Iterable[Iterable[str]]) -> None:
    headers = list(headers)
    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = str(header)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)


def add_bullet(document: Document, text: str) -> None:
    try:
        document.add_paragraph(text, style="List Bullet")
    except KeyError:
        document.add_paragraph(f"- {text}")


def append_generation_preprocess_section(
    doc: Document,
    summary: Dict,
    style_summary: Dict,
    subset_summary: Dict,
    benchmark_summary: Dict,
    raw_bytes: int,
    processed_bytes: int,
) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("생성모델 전처리 결과 보강본", level=1)
    doc.add_paragraph(
        "기존 세그멘테이션 양식을 기반으로, 생성모델 학습에 직접 사용할 수 있는 전처리 결과와 고정 벤치마크 구성을 추가로 정리했다."
    )

    add_table(
        doc,
        ["항목", "값"],
        [
            ["원본 이미지 수", f"{summary['total_images_seen']:,}"],
            ["채택된 생성 학습 샘플", f"{summary['accepted_recon_samples']:,}"],
            ["헤어 관련 prompt bank", f"{summary['hair_related_prompt_records']:,}"],
            ["학습 train_all", f"{subset_summary['train_all']:,}"],
            ["학습 train_highconf", f"{subset_summary['train_highconf']:,}"],
            ["평가 val/eval", f"{subset_summary['eval_rows']:,}"],
            ["원본 데이터 크기", format_bytes(raw_bytes)],
            ["전처리 데이터 크기", format_bytes(processed_bytes)],
        ],
    )

    doc.add_paragraph("스타일 pseudo-label 분포")
    top_style_rows = [
        [style_name, str(count)]
        for style_name, count in list(style_summary["top_style_counts"].items())[:10]
    ]
    add_table(doc, ["스타일", "샘플 수"], top_style_rows)

    doc.add_paragraph("테스트 데이터셋 구성")
    add_table(
        doc,
        ["셋", "샘플 수", "용도"],
        [
            ["smoke_eval_quick", str(benchmark_summary.get("smoke_rows", 0)), "10분 내 smoke test / cold-start sanity check"],
            ["recon_eval_stratified", str(benchmark_summary.get("recon_rows", 0)), "hair 영역 정량 복원 평가"],
            ["edit_eval_requests", str(benchmark_summary.get("edit_rows", 0)), "헤어 관련 요청 prompt-following 정성 평가"],
        ],
    )


def append_training_section(
    doc: Document,
    stage1_config: Dict,
    stage2_config: Dict,
    subset_summary: Dict,
) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("생성모델 학습 계획 보강본", level=1)
    doc.add_paragraph(
        "세그멘테이션 재학습은 제외하고, 생성모델만 UNet LoRA 방식으로 미세조정한다. 목적은 헤어스타일 생성 성능을 높이면서 얼굴과 의상 보존을 유지하는 것이다."
    )
    add_table(
        doc,
        ["구분", "설정"],
        [
            ["Base model", stage1_config["pretrained_model_name_or_path"]],
            ["ControlNet", stage1_config["controlnet_model_name_or_path"]],
            ["Trainable weights", "UNet LoRA only"],
            ["Train samples", f"{subset_summary['train_all']:,}"],
            ["Validation samples", "96 fixed recon benchmark"],
            ["Precision", stage1_config["mixed_precision"]],
            ["Optimizer", "8-bit AdamW" if stage1_config["use_8bit_adam"] else "AdamW"],
        ],
    )

    doc.add_paragraph("학습 stage 구성")
    add_table(
        doc,
        ["Stage", "Manifest", "Steps", "LR", "Dropout", "메모"],
        [
            [
                "Stage 1",
                Path(stage1_config["manifest"]).name,
                str(stage1_config["max_train_steps"]),
                str(stage1_config["learning_rate"]),
                str(stage1_config["lora_dropout"]),
                "넓은 style coverage 확보",
            ],
            [
                "Stage 2",
                Path(stage2_config["manifest"]).name,
                str(stage2_config["max_train_steps"]),
                str(stage2_config["learning_rate"]),
                str(stage2_config["lora_dropout"]),
                "high-confidence style sharpen",
            ],
        ],
    )

    doc.add_paragraph("과적합 방지 기법")
    for bullet in [
        "LoRA only fine-tuning",
        "prompt / face / control dropout",
        "mask jitter",
        "min-SNR weighting",
        "weighted sampling with style confidence",
        "validation every 250 steps and patience 3 early stopping",
    ]:
        add_bullet(doc, bullet)


def append_model_section(doc: Document, subset_summary: Dict, benchmark_summary: Dict) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("학습 산출물 계획 보강본", level=1)
    doc.add_paragraph(
        "본 문서는 실제 학습 실행 전 준비 상태를 기록한 초안이다. 체크포인트 생성 후 best/final 경로와 정량지표를 다시 채워 제출본으로 확정한다."
    )
    add_table(
        doc,
        ["항목", "현재 상태"],
        [
            ["학습 데이터 수", f"{subset_summary['train_all']:,}"],
            ["고신뢰 stage 2 데이터 수", f"{subset_summary['train_highconf']:,}"],
            ["평가 데이터 수", f"{subset_summary['eval_rows']:,}"],
            ["smoke 테스트 수", str(benchmark_summary.get("smoke_rows", 0))],
            ["정량 recon 테스트 수", str(benchmark_summary.get("recon_rows", 0))],
            ["정성 edit 테스트 수", str(benchmark_summary.get("edit_rows", 0))],
            ["체크포인트 상태", "학습 실행 전"],
            ["최종 배포 방식", "기존 serverless endpoint에 runtime LoRA 로딩"],
        ],
    )

    doc.add_paragraph("실행 후 채워야 할 항목")
    for bullet in [
        "stage1 best / final 경로",
        "stage2 best / final 경로",
        "hair_mae / hair_psnr / hair_ssim",
        "face_preserve_mae / cloth_preserve_mae",
        "stage1 vs stage2 qualitative 비교 결과",
    ]:
        add_bullet(doc, bullet)


def build_preprocess_doc(
    template_path: Path,
    output_path: Path,
    summary: Dict,
    style_summary: Dict,
    subset_summary: Dict,
    benchmark_summary: Dict,
    raw_bytes: int,
    processed_bytes: int,
) -> None:
    doc = Document(template_path)
    replace_paragraph_prefixes(
        doc,
        {
            "프로젝트명:": "프로젝트명: MirrAI - 헤어스타일 생성모델 학습용 전처리 및 벤치마크 구축",
            "전처리 목적:": "전처리 목적: CelebA-Dialog HQ를 기반으로 헤어스타일 생성모델 학습에 바로 투입할 수 있는 이미지, 마스크, 조건 입력, 평가셋을 구축한다.",
            "문제 정의:": "문제 정의: 생성모델 학습에는 source/target 이미지, hair mask, control image, face crop, 고정 validation/test 세트가 함께 필요하며, pseudo-style 라벨 신뢰도까지 관리해야 한다.",
            "추가 고려 사항:": "추가 고려 사항: 4090 24GB / 50GB volume 제약 안에서 학습 가능한 크기로 유지하고, 과적합 감시용 smoke/recon/edit benchmark를 동시에 만든다.",
            "데이터 출처 및 수집 방법:": "데이터 출처 및 수집 방법: CelebA-Dialog HQ 공식 데이터 수집 후 로컬 전처리 스크립트로 generation manifest 생성",
            "최종 학습 직접 사용 범위:": "최종 학습 직접 사용 범위: 전처리된 generation manifest와 benchmark manifest만 직접 학습/평가에 사용",
            "원본 데이터 샘플 형태:": "원본 데이터 샘플 형태: 인물 이미지, attribute caption, request annotation. 전처리 후 target/source image, hair mask, canny, face crop, pseudo style label을 포함",
            "전체 흐름도:": "전체 흐름도: 원본 다운로드 -> 품질 필터 -> hair/face/cloth mask 정리 -> canny/face crop 생성 -> pseudo-style enrichment -> curriculum manifest 저장 -> smoke/recon/edit benchmark 저장",
        },
    )
    append_generation_preprocess_section(
        doc,
        summary=summary,
        style_summary=style_summary,
        subset_summary=subset_summary,
        benchmark_summary=benchmark_summary,
        raw_bytes=raw_bytes,
        processed_bytes=processed_bytes,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def build_training_doc(
    template_path: Path,
    output_path: Path,
    stage1_config: Dict,
    stage2_config: Dict,
    subset_summary: Dict,
) -> None:
    doc = Document(template_path)
    replace_paragraph_prefixes(
        doc,
        {
            "AIHub 한국인 헤어스타일 이미지 HQ 데이터셋으로 hair-only segmentation 모델을 학습하고 결과를 정리한다.": "CelebA-Dialog HQ 기반 generation manifest로 헤어스타일 생성 LoRA 학습을 준비하고 결과를 정리한다.",
            "비교 대상 모델 수:": "비교 대상 모델 수: 총 2종(stage1 curriculum / stage2 high-confidence refinement)",
            "최종 선정 모델:": "최종 선정 모델: SD 1.5 Inpainting + ControlNet(Canny) + IP-Adapter(face) + UNet LoRA",
            "아키텍처 개요:": "아키텍처 개요: source image, hair mask, canny control, face crop을 입력으로 받아 hair 영역만 자연스럽게 재생성하도록 UNet LoRA를 학습한다.",
            "입력 이미지 -> backbone encoder -> segmentation decoder -> hair class logit 추출 -> binary hair mask 예측": "source image + mask + canny + face crop -> frozen base pipeline -> trainable UNet LoRA -> hair-aware inpaint output",
            "학습 데이터 메모:": "학습 데이터 메모: 전처리된 CelebA-Dialog HQ generation set(train_all 25,288 / highconf 15,629 / eval 1,167)을 사용한다.",
            "인프라 메모:": "인프라 메모: RunPod RTX 4090 24GB 환경에서 50GB volume을 사용하고, serverless는 smoke/eval 전용으로 분리한다.",
        },
    )
    append_training_section(doc, stage1_config=stage1_config, stage2_config=stage2_config, subset_summary=subset_summary)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def build_model_doc(
    template_path: Path,
    output_path: Path,
    subset_summary: Dict,
    benchmark_summary: Dict,
) -> None:
    doc = Document(template_path)
    replace_paragraph_prefixes(
        doc,
        {
            "프로젝트명:": "프로젝트명: MirrAI - 헤어스타일 생성모델 UNet LoRA",
            "문제 정의:": "문제 정의: 얼굴과 의상은 유지하면서 hair 영역만 스타일 조건에 맞게 자연스럽게 다시 생성하는 모델을 구축한다.",
            "모델 목적:": "모델 목적: 기존 마스킹 파이프라인과 결합해 identity-preserving hairstyle generation 성능을 높인다.",
            "선정 모델:": "선정 모델: SD 1.5 Inpainting + ControlNet(Canny) + IP-Adapter(face) + UNet LoRA",
            "아키텍처 개요:": "아키텍처 개요: frozen base diffusion pipeline 위에 UNet LoRA만 얹어 budget-friendly fine-tuning을 수행한다.",
            "아키텍처 시각화 또는 참고 파일:": "아키텍처 시각화 또는 참고 파일: docs/generation_training_runbook.md 및 scripts/train_hair_lora.py 참고",
            "설계 근거:": "설계 근거: full fine-tuning 대신 LoRA만 학습하면 4090 24GB 예산 환경에서도 과적합과 비용을 동시에 제어할 수 있다.",
            "학습 데이터 수:": f"학습 데이터 수: {subset_summary['train_all']:,}건",
            "검증 데이터 수:": "검증 데이터 수: 96건(fixed recon benchmark)",
            "평가 데이터 수:": f"평가 데이터 수: {benchmark_summary.get('edit_rows', 0) + benchmark_summary.get('recon_rows', 0)}건",
            "완료 epoch 수:": "완료 epoch 수: 학습 실행 후 갱신",
            "평균 epoch 시간(초):": "평균 epoch 시간(초): 학습 실행 후 갱신",
        },
    )
    append_model_section(doc, subset_summary=subset_summary, benchmark_summary=benchmark_summary)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    summary = load_json(REPORTS_DIR / "summary.json")
    style_summary = load_json(REPORTS_DIR / "style_enrichment_summary.json")
    subset_summary = load_json(REPORTS_DIR / "style_training_subsets_summary.json")
    benchmark_path = REPORTS_DIR / "benchmark_summary.json"
    benchmark_summary = load_json(benchmark_path) if benchmark_path.exists() else {"smoke_rows": 0, "recon_rows": 0, "edit_rows": 0}
    stage1_config = load_json(CONFIG_DIR / "stage1_budget_4090.json")
    stage2_config = load_json(CONFIG_DIR / "stage2_highconf_4090.json")

    raw_bytes = compute_dir_size(REPO_ROOT / "dataset_build" / "raw" / "celeba_dialog_hq")
    processed_bytes = compute_dir_size(REPO_ROOT / "dataset_build" / "processed" / "celeba_dialog_hq_generation")

    build_preprocess_doc(
        DOCS_DIR / "MirrAI_데이터 전처리_인공지능 데이터 전처리 결과서.docx.docx",
        OUTPUT_DIR / "MirrAI_생성모델_데이터 전처리 결과서.docx",
        summary=summary,
        style_summary=style_summary,
        subset_summary=subset_summary,
        benchmark_summary=benchmark_summary,
        raw_bytes=raw_bytes,
        processed_bytes=processed_bytes,
    )
    build_training_doc(
        DOCS_DIR / "MirrAI_데이터 전처리_인공지능 학습 결과서_22기.docx.docx",
        OUTPUT_DIR / "MirrAI_생성모델_인공지능 학습 결과서.docx",
        stage1_config=stage1_config,
        stage2_config=stage2_config,
        subset_summary=subset_summary,
    )
    build_model_doc(
        DOCS_DIR / "MirrAI_데이터 전처리_학습된 인공지능 모델_22기.docx.docx",
        OUTPUT_DIR / "MirrAI_생성모델_학습된 인공지능 모델.docx",
        subset_summary=subset_summary,
        benchmark_summary=benchmark_summary,
    )

    print(json.dumps({"output_dir": str(OUTPUT_DIR)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
